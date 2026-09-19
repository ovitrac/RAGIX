"""Legacy Office conversion gates, using only constructed documents.

Author: Olivier Vitrac, PhD, HDR | olivier.vitrac@adservio.fr | Adservio
"""

from dataclasses import replace
import hashlib
from pathlib import Path
import shutil
import sys
from urllib.parse import unquote, urlparse
import zipfile

import pytest

from ragix_kernels.base import KernelInput
from ragix_kernels.saqqara import office_conversion as office
from ragix_kernels.saqqara.adapters import (
    ConversionUnavailable,
    UnreadableFile,
    adapter_for,
    read_corpus,
    read_path,
)
from ragix_kernels.saqqara.assets import AssetStore
from ragix_kernels.saqqara.builder import build_tree
from ragix_kernels.saqqara.kernels.saqqara_run import SaqqaraKernel
from ragix_kernels.saqqara.model import Tree


def native(path):
    if path.suffix == ".docx":
        from docx import Document

        doc = Document()
        doc.add_heading("Synthetic document", 1)
        doc.add_paragraph("A paragraph with 23 mm.")
        table = doc.add_table(rows=3, cols=2)
        for row, values in zip(
            table.rows, [("Name", "Value"), ("Length", "23 mm"), ("Width", "17 mm")]
        ):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.save(path)
    else:
        from openpyxl import Workbook

        book = Workbook()
        sheet = book.active
        sheet.title = "Synthetic"
        sheet.append(["Name", "Value"])
        sheet.append(["Length", 23])
        sheet["B3"] = "=B2*2"
        sheet.merge_cells("A4:B4")
        sheet["A4"] = "Merged title"
        book.create_sheet("Hidden").sheet_state = "hidden"
        book.save(path)
    return path


@pytest.fixture
def fake_converter(monkeypatch):
    calls = []
    monkeypatch.setattr(office, "libreoffice_executable", lambda: "/fake/libreoffice")

    def run(args, **kwargs):
        if "--version" in args:
            return 0, "LibreOffice test-build\n"
        profile = Path(unquote(urlparse(args[1].split("=", 1)[1]).path))
        assert (profile / "user" / "registrymodifications.xcu").read_text() == office.PROFILE
        assert "--headless" in args
        source = Path(args[-1])
        assert source.name in ("source.doc", "source.xls")
        output = Path(args[args.index("--outdir") + 1])
        target = args[args.index("--convert-to") + 1].split(":")[0]
        native(output / ("source." + target))
        calls.append((profile.parent, source.read_bytes()))
        return 0, ""

    monkeypatch.setattr(office, "_run", run)
    return calls


@pytest.mark.parametrize("suffix,target", [(".doc", "docx"), (".XLS", "xlsx")])
def test_conversion_uses_native_reader_and_bridges_provenance(
    tmp_path, fake_converter, suffix, target
):
    source = tmp_path / ("Original with spaces" + suffix)
    source.write_bytes(b"synthetic legacy input")
    store = AssetStore(tmp_path / "assets")
    records = read_path(source, conversion_store=store)
    bridge = records[0].conversion
    assert bridge.input_format == suffix.lower()[1:]
    assert bridge.output_format == target
    assert bridge.source_sha256 == hashlib.sha256(source.read_bytes()).hexdigest()
    assert bridge.artifact_retained
    derived = tmp_path / ("retained." + target)
    derived.write_bytes(store.read(bridge.derived_sha256))
    assert [replace(r, conversion=None).to_dict() for r in records] == [
        r.to_dict() for r in read_path(derived)
    ]
    assert source.read_bytes() == b"synthetic legacy input"
    assert not fake_converter[0][0].exists()
    assert fake_converter[0][1] == source.read_bytes()
    adapter = adapter_for(source)
    tree = build_tree(records, str(source), target, target, adapter.version).tree
    assert tree.meta["conversion"] == bridge.to_dict()
    assert Tree.from_dict(tree.to_dict()).to_dict() == tree.to_dict()
    for node in tree.walk():
        assert node.provenance.chain[0] == bridge
        assert node.provenance.source_sha256 == bridge.source_sha256
        assert node.provenance.source_path == str(source)
        assert node.provenance.leaf.format != "conversion"
    with pytest.raises(ValueError, match="inconsistent"):
        build_tree(
            [records[0], replace(records[-1], conversion=None)],
            str(source),
            target,
            target,
            adapter.version,
        )


@pytest.mark.parametrize("suffix", [".doc", ".xls"])
def test_missing_libreoffice_is_a_counted_refusal(tmp_path, monkeypatch, suffix):
    source = tmp_path / ("legacy" + suffix)
    source.write_bytes(b"synthetic")
    monkeypatch.setattr(office.shutil, "which", lambda name: None)
    with pytest.raises(ConversionUnavailable, match="Install LibreOffice"):
        read_path(source)
    records, report = read_corpus([source])
    assert not records and not report.read
    assert len(report.refusals) == 1
    assert report.refusals[0].reason == "converter-unavailable"
    assert "Install LibreOffice" in report.refusals[0].detail
    result = SaqqaraKernel().compute(
        KernelInput(tmp_path / "workspace", {"source": {"path": str(source)}})
    )
    assert result["documents"] == []
    assert result["report"]["refusals"][0]["reason"] == "converter-unavailable"


@pytest.mark.parametrize("suffix", [".docx", ".xlsx"])
def test_native_formats_do_not_need_libreoffice(tmp_path, monkeypatch, suffix):
    monkeypatch.setattr(office.shutil, "which", lambda name: None)
    records = read_path(native(tmp_path / ("native" + suffix)))
    assert records and all("conversion" not in r.to_dict() for r in records)


@pytest.mark.parametrize(
    "failure",
    [
        "exit",
        "absent",
        "invalid",
        "wrong-format",
        "oversize",
        "expanded",
        "traversal",
        "changed-source",
    ],
)
def test_bad_conversion_never_returns_partial_records(
    tmp_path, monkeypatch, fake_converter, failure
):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"synthetic")
    run = office._run

    def broken(args, **kwargs):
        result = run(args, **kwargs)
        if "--version" in args:
            return result
        output = Path(args[args.index("--outdir") + 1]) / "source.docx"
        if failure == "exit":
            return 1, "failure"
        if failure == "absent":
            output.unlink()
        elif failure == "invalid":
            output.write_bytes(b"not zip")
        elif failure == "wrong-format":
            native(output.with_suffix(".xlsx"))
            output.write_bytes(output.with_suffix(".xlsx").read_bytes())
        elif failure == "traversal":
            with zipfile.ZipFile(output, "a") as archive:
                archive.writestr("../escape", "invalid")
        elif failure == "changed-source":
            source.write_bytes(b"changed")
        return result

    monkeypatch.setattr(office, "_run", broken)
    options = office.ConversionOptions(
        max_output_bytes=1 if failure == "oversize" else 128 * 1024 * 1024,
        max_expanded_bytes=1 if failure == "expanded" else 512 * 1024 * 1024,
    )
    with pytest.raises(UnreadableFile):
        list(office.read_legacy(source, adapter_for(source)._read_native, options=options))
    assert not fake_converter[0][0].exists()


def test_timeout_kills_process_and_refuses():
    with pytest.raises(UnreadableFile, match="timed out"):
        office._run([sys.executable, "-c", "import time; time.sleep(20)"], timeout=0.05)


def test_input_limit_and_empty_reader_refuse(tmp_path, fake_converter):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"synthetic")
    with pytest.raises(UnreadableFile, match="input exceeds"):
        list(
            office.read_legacy(
                source, lambda path: (), options=office.ConversionOptions(max_input_bytes=1)
            )
        )
    with pytest.raises(UnreadableFile, match="no readable observations"):
        list(office.read_legacy(source, lambda path: ()))


@pytest.mark.parametrize("suffix", [".doc", ".xls"])
def test_kernel_retains_derivative_and_runs_analyzers(tmp_path, fake_converter, suffix):
    source = tmp_path / ("legacy" + suffix)
    source.write_bytes(b"synthetic")
    workspace = tmp_path / "workspace"
    result = SaqqaraKernel().compute(KernelInput(workspace, {"source": {"path": str(source)}}))
    assert len(result["documents"]) == 1
    document = result["documents"][0]
    conversion = document["tree"]["meta"]["conversion"]
    assert conversion["artifact_retained"]
    assert document["sha256"] == conversion["source_sha256"]
    assert AssetStore(workspace / "assets" / "office-conversions").read(
        conversion["derived_sha256"]
    )
    assert "builder" in document["traces"] and len(document["traces"]) > 1


@pytest.mark.skipif(
    not (shutil.which("libreoffice") or shutil.which("soffice")),
    reason="optional LibreOffice is absent",
)
@pytest.mark.parametrize(
    "target,legacy,filter_name", [("docx", "doc", "MS Word 97"), ("xlsx", "xls", "MS Excel 97")]
)
def test_real_binary_conversion(tmp_path, target, legacy, filter_name):
    if legacy == "doc":
        # Construct the binary fixture from RTF: direct DOCX -> DOC export in
        # some LibreOffice versions flattens the table before this reader runs.
        seed = tmp_path / "seed.rtf"
        seed.write_text(
            r"{\rtf1\ansi Synthetic document\par A paragraph with 23 mm.\par "
            r"\trowd\cellx2000\cellx4000\intbl Name\cell Value\cell\row "
            r"\trowd\cellx2000\cellx4000\intbl Length\cell 23 mm\cell\row "
            r"\trowd\cellx2000\cellx4000\intbl Width\cell 17 mm\cell\row "
            r"\pard\par}",
            encoding="ascii",
        )
    else:
        seed = native(tmp_path / ("seed." + target))
    output = tmp_path / "binary"
    output.mkdir()
    status, _ = office._run(
        [
            office.libreoffice_executable(),
            "-env:UserInstallation=" + (tmp_path / "seed-profile").as_uri(),
            "--headless",
            "--convert-to",
            legacy + ":" + filter_name,
            "--outdir",
            str(output),
            str(seed),
        ],
        timeout=120,
    )
    assert status == 0
    source = output / ("seed." + legacy)
    original = source.read_bytes()
    assert original[:8] == bytes.fromhex("d0cf11e0a1b11ae1")
    store = AssetStore(tmp_path / "retained")
    records = read_path(source, conversion_store=store)
    assert records and source.read_bytes() == original
    assert records[0].conversion.converter_version.startswith("LibreOffice ")
    texts = [r.text for r in records]
    if legacy == "doc":
        assert "A paragraph with 23 mm." in texts
        assert "23 mm" in texts and "17 mm" in texts
    else:
        assert any(r.facts.get("formula") == "=B2*2" for r in records)
        assert any(r.facts.get("hidden") for r in records)
        assert "Merged title" in texts
    derived = tmp_path / ("derived." + target)
    derived.write_bytes(store.read(records[0].conversion.derived_sha256))
    assert [replace(r, conversion=None).to_dict() for r in records] == [
        r.to_dict() for r in read_path(derived)
    ]
