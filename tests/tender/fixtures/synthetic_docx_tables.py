"""Synthetic docx tables — ground truth for the docx table lane (PD0).

Author: Olivier Vitrac, PhD, HDR | olivier.vitrac@adservio.fr | Adservio | 2026-09-14

Fixtures MHD1-MHD7 per ADJUDICATION_DESIGN_DOCX_TABLES_20260730 §5 (PD0), with
INVENTED content only — no client text, no personal data. ``build`` writes the
main document (MHD1-MHD6: body, page-header and nested tables, layout tables,
fillable markers in and out of cells); ``build_twin`` writes the MHD7
blank/filled pair whose slot arithmetic mirrors the DC1 measurement
(F-DX3: empty(blank) - empty(filled) = answers).

``EXPECTED`` addresses each table by its LOCATION STREAM (adjudication L-R1:
``table_index`` counted within body / header:<section> / nested), the shape
the future docx locator freezes. The integrity gate
``tests/test_gate_gxd0_docx_fixture.py`` keeps this file honest by re-reading
the raw OOXML zip — it never uses python-docx nor any substrate kernel.

The F-DX4 lesson is baked in: FORMTEXT and ``w:sdt`` sit INSIDE table cells
(the marker-as-slot case), FORMCHECKBOX sits in a body paragraph OUTSIDE any
table (the DC1 reality) — a correct Layer A must tell the two apart.

``n_empty`` counts EXCLUDE vMerge continuation cells: a continuation tc has
no text by construction (its value lives at the merge anchor) and is a
structural blank, never an answer slot. Naive text-empty counting overcounts
on merged tables — Layer A must project anchors before any slot arithmetic.
"""

from __future__ import annotations

from pathlib import Path

MHD1 = "MHD1 deux-tiers"      # two-tier column header: gridSpan tier + vMerge across tiers
MHD2 = "MHD2 tuiles label"    # vMerge label-column tiling + empty slot column
MHD3A = "MHD3a layout 1x1"    # layout table, prose, no style/borders
MHD3B = "MHD3b layout 1x2"    # layout table, two prose cells
MHD4 = "MHD4 marqueurs"       # FORMTEXT + w:sdt in cells, plain slot; FORMCHECKBOX outside
MHD5 = "MHD5 en-tete page"    # table living in the page header (F-DX5)
MHD6 = "MHD6 imbrique"        # nested table inside a cell (F-DX5)
MHD7 = "MHD7 jumeaux"         # blank/filled twin pair (F-DX3 slot arithmetic)

SHADE = "D9D9D9"              # header shading fill (A-R2 style fact)

EXPECTED = {
    MHD1: {
        "location": "body", "table_index": 0, "style": "TableGrid",
        "n_rows": 4, "n_grid_cols": 4,
        # raw tc count per row AFTER python-docx horizontal merge (tcs removed)
        "tc_counts": [2, 4, 4, 4],
        "grid_spans": [{"row": 0, "tc": 1, "span": 3}],          # "Moyens mobilisés"
        "vmerge_restart": [{"row": 0, "tc": 0}],                 # "Lot" spans both tiers
        "vmerge_continue": [{"row": 1, "tc": 0}],
        "texts": {(0, 0): "Lot", (0, 1): "Moyens mobilisés",
                  (1, 1): "Humains", (1, 2): "Techniques", (1, 3): "Total",
                  (2, 0): "Lot A", (3, 0): "Lot B"},
        "bold": [(0, 0), (0, 1), (1, 1), (1, 2), (1, 3)],
        "shaded": [(0, 0), (0, 1)],
        "n_empty": 0,
    },
    MHD2: {
        "location": "body", "table_index": 1, "style": "TableGrid",
        "n_rows": 4, "n_grid_cols": 3,
        "tc_counts": [3, 3, 3, 3],
        "grid_spans": [],
        "vmerge_restart": [{"row": 0, "tc": 0}, {"row": 2, "tc": 0}],
        "vmerge_continue": [{"row": 1, "tc": 0}, {"row": 3, "tc": 0}],
        "texts": {(0, 0): "Phase amont", (2, 0): "Phase aval",
                  (0, 1): "cadrer", (1, 1): "planifier",
                  (2, 1): "produire", (3, 1): "livrer"},
        "bold": [(0, 0), (2, 0)],
        "shaded": [],
        "n_empty": 4,                                            # the whole slot column
        "slot_tcs": [(0, 2), (1, 2), (2, 2), (3, 2)],
    },
    MHD3A: {
        "location": "body", "table_index": 2, "style": None,
        "n_rows": 1, "n_grid_cols": 1, "tc_counts": [1],
        "grid_spans": [], "vmerge_restart": [], "vmerge_continue": [],
        "texts": {(0, 0): "Ce cadre présente le contexte général de la consultation fictive."},
        "bold": [], "shaded": [], "n_empty": 0,
    },
    MHD3B: {
        "location": "body", "table_index": 3, "style": None,
        "n_rows": 1, "n_grid_cols": 2, "tc_counts": [2],
        "grid_spans": [], "vmerge_restart": [], "vmerge_continue": [],
        "texts": {(0, 0): "Colonne de gauche, texte courant.",
                  (0, 1): "Colonne de droite, texte courant."},
        "bold": [], "shaded": [], "n_empty": 0,
    },
    MHD4: {
        "location": "body", "table_index": 4, "style": "TableGrid",
        "n_rows": 3, "n_grid_cols": 2, "tc_counts": [2, 2, 2],
        "grid_spans": [], "vmerge_restart": [], "vmerge_continue": [],
        "texts": {(0, 0): "Nom du répondant", (1, 0): "Ville", (2, 0): "Observations"},
        "bold": [], "shaded": [],
        "n_empty": 3,                                            # marker cells carry no text
        "formtext_tc": (0, 1),                                   # legacy field IN a cell
        "sdt_tc": (1, 1),                                        # content control IN a cell
        "plain_slot_tc": (2, 1),                                 # empty, no marker at all
    },
    MHD5: {
        "location": "header:0", "table_index": 0, "style": None,
        "n_rows": 1, "n_grid_cols": 2, "tc_counts": [2],
        "grid_spans": [], "vmerge_restart": [], "vmerge_continue": [],
        "texts": {(0, 0): "Référence du dossier", (0, 1): "OFF-2026-014"},
        "bold": [], "shaded": [], "n_empty": 0,
    },
    MHD6: {
        "location": "body", "table_index": 5, "style": None,
        "n_rows": 1, "n_grid_cols": 2, "tc_counts": [2],
        "grid_spans": [], "vmerge_restart": [], "vmerge_continue": [],
        "texts": {(0, 0): "Détails techniques"},
        "bold": [], "shaded": [], "n_empty": 0,
        "nested": {"parent_tc": (0, 1), "n_rows": 2, "n_grid_cols": 2,
                   "texts": {(0, 0): "Débit", (0, 1): "10",
                             (1, 0): "Latence", (1, 1): "2"}},
    },
    MHD7: {
        "location": "body", "table_index": 0, "style": "TableGrid",
        "n_rows": 3, "n_grid_cols": 3, "tc_counts": [3, 3, 3],
        "texts": {(0, 0): "N°", (0, 1): "Désignation", (0, 2): "Montant",
                  (1, 0): "1", (2, 0): "2"},
        "bold": [(0, 0), (0, 1), (0, 2)],
        "blank_empty": 4,                                        # (1,1)(1,2)(2,1)(2,2)
        "filled_empty": 2,                                       # (2,1)(2,2) stay blank
        "answers": [(1, 1), (1, 2)],                             # filled - blank = 2
        "answer_texts": {(1, 1): "Prestation d'étude fictive", (1, 2): "1200"},
    },
}

# document-level markers OUTSIDE tables (the F-DX4 reality on DC1)
BODY_CHECKBOX_TEXT = "j'atteste l'exactitude des renseignements fournis"


# --------------------------------------------------------------------- helpers

def _qn(tag):
    from docx.oxml.ns import qn
    return qn(tag)

def _el(tag, **attrs):
    from docx.oxml import OxmlElement
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(_qn(k), v)
    return e


def _add_field(paragraph, instr: str, ff_child: str, name: str) -> None:
    """Append a legacy form field (FORMTEXT / FORMCHECKBOX) to a paragraph."""
    p = paragraph._p
    r1 = _el("w:r")
    fld = _el("w:fldChar", **{"w:fldCharType": "begin"})
    ff = _el("w:ffData")
    ff.append(_el("w:name", **{"w:val": name}))
    ff.append(_el("w:enabled"))
    ff.append(_el(ff_child))
    fld.append(ff)
    r1.append(fld)
    r2 = _el("w:r")
    it = _el("w:instrText")
    it.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    it.text = f" {instr} "
    r2.append(it)
    r3 = _el("w:r")
    r3.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    r4 = _el("w:r")
    r4.append(_el("w:fldChar", **{"w:fldCharType": "end"}))
    for r in (r1, r2, r3, r4):
        p.append(r)


def _add_sdt(cell, alias: str) -> None:
    """Append an empty block-level content control to a table cell."""
    sdt = _el("w:sdt")
    pr = _el("w:sdtPr")
    pr.append(_el("w:alias", **{"w:val": alias}))
    sdt.append(pr)
    content = _el("w:sdtContent")
    content.append(_el("w:p"))
    sdt.append(content)
    cell._tc.append(sdt)


def _put(cell, text: str, bold: bool = False, shade: bool = False) -> None:
    par = cell.paragraphs[0]
    run = par.add_run(text)
    run.bold = bold
    if shade:
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(_el("w:shd", **{"w:val": "clear", "w:fill": SHADE}))


# --------------------------------------------------------------------- builders

def build(path: Path) -> Path:
    """Write the main synthetic document (MHD1-MHD6) at ``path``."""
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    doc.add_paragraph("Document synthétique — tableaux docx (contenu inventé).")

    # ---- MHD1: two-tier column header (gridSpan tier + vMerge across tiers)
    t = doc.add_table(rows=4, cols=4, style="Table Grid")
    t.cell(0, 0).merge(t.cell(1, 0))
    _put(t.cell(0, 0), "Lot", bold=True, shade=True)
    t.cell(0, 1).merge(t.cell(0, 3))
    _put(t.cell(0, 1), "Moyens mobilisés", bold=True, shade=True)
    for c, txt in [(1, "Humains"), (2, "Techniques"), (3, "Total")]:
        _put(t.cell(1, c), txt, bold=True)
    for r, (lot, vals) in enumerate([("Lot A", ["4", "2", "6"]),
                                     ("Lot B", ["3", "1", "4"])], start=2):
        _put(t.cell(r, 0), lot)
        for c, v in enumerate(vals, start=1):
            _put(t.cell(r, c), v)

    doc.add_paragraph("Répartition des phases (contenu inventé).")

    # ---- MHD2: vMerge label tiling + slot column
    t = doc.add_table(rows=4, cols=3, style="Table Grid")
    t.cell(0, 0).merge(t.cell(1, 0))
    _put(t.cell(0, 0), "Phase amont", bold=True)
    t.cell(2, 0).merge(t.cell(3, 0))
    _put(t.cell(2, 0), "Phase aval", bold=True)
    for r, txt in enumerate(["cadrer", "planifier", "produire", "livrer"]):
        _put(t.cell(r, 1), txt)
    # column 2 stays empty on purpose: the slot column

    # ---- MHD3a / MHD3b: layout tables (no style, prose)
    t = doc.add_table(rows=1, cols=1)
    _put(t.cell(0, 0), "Ce cadre présente le contexte général de la consultation fictive.")
    t = doc.add_table(rows=1, cols=2)
    _put(t.cell(0, 0), "Colonne de gauche, texte courant.")
    _put(t.cell(0, 1), "Colonne de droite, texte courant.")

    # ---- MHD4: fillable markers IN cells + a plain slot
    t = doc.add_table(rows=3, cols=2, style="Table Grid")
    _put(t.cell(0, 0), "Nom du répondant")
    _add_field(t.cell(0, 1).paragraphs[0], "FORMTEXT", "w:textInput", "Texte1")
    _put(t.cell(1, 0), "Ville")
    _add_sdt(t.cell(1, 1), "Ville")
    _put(t.cell(2, 0), "Observations")
    # cell (2,1): plain empty slot, no marker

    # FORMCHECKBOX in a body paragraph OUTSIDE any table (F-DX4)
    p = doc.add_paragraph("Engagement : ")
    _add_field(p, "FORMCHECKBOX", "w:checkBox", "Case1")
    p.add_run(" " + BODY_CHECKBOX_TEXT + ".")

    # ---- MHD6: nested table inside a cell
    t = doc.add_table(rows=1, cols=2)
    _put(t.cell(0, 0), "Détails techniques")
    inner = t.cell(0, 1).add_table(rows=2, cols=2)
    for (r, c), txt in EXPECTED[MHD6]["nested"]["texts"].items():
        _put(inner.cell(r, c), txt)

    # ---- MHD5: table in the page header
    header = doc.sections[0].header
    header.is_linked_to_previous = False
    ht = header.add_table(rows=1, cols=2, width=Cm(12))
    _put(ht.cell(0, 0), "Référence du dossier")
    _put(ht.cell(0, 1), "OFF-2026-014")

    doc.save(path)
    return path


def build_twin(blank_path: Path, filled_path: Path) -> tuple[Path, Path]:
    """Write the MHD7 blank/filled pair — identical structure, two answers added."""
    from docx import Document

    for path, filled in [(blank_path, False), (filled_path, True)]:
        doc = Document()
        doc.add_paragraph("Bordereau synthétique (contenu inventé).")
        t = doc.add_table(rows=3, cols=3, style="Table Grid")
        for c, txt in enumerate(["N°", "Désignation", "Montant"]):
            _put(t.cell(0, c), txt, bold=True)
        _put(t.cell(1, 0), "1")
        _put(t.cell(2, 0), "2")
        if filled:
            for (r, c), txt in EXPECTED[MHD7]["answer_texts"].items():
                _put(t.cell(r, c), txt)
        doc.save(path)
    return blank_path, filled_path
