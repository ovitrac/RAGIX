"""
saqqara.adapters.pptx — the presentation reader.

Author: Olivier Vitrac, PhD, HDR | olivier.vitrac@adservio.fr | Adservio | 2026-08-27

Carries K2.15 of SPEC.md.

Slides are numbered the way a reader counts them, from one. A zero-based slide number is correct
in the file and wrong in a citation: nobody looking for slide 7 wants slide 8.

**A table on a slide is a table.** It arrives as a shape with no text frame, so a reader that
walks text frames sees nothing at all — not an empty table, but no table. Its cells are emitted
here in the same vocabulary a word-processing or spreadsheet grid is emitted in, which means the
header and label rules read a slide table without learning that slides exist.

Speaker notes carry their own address (`notes=True`) as well as their own fact
(`on_slide=False`). The two are not redundant: the address says where the text is, the fact says
what it is. Speaker notes are kept apart from what is on the slide. They are frequently the most substantive
text in a deck, and they are also the text the audience never saw — merging the two produces a
document that claims to show what it merely mentioned.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterator

from ..model import PptxLocator
from .contract import (
    FIGURE_FACTS,
    PART_SKIPS,
    GRID_CELL_FACTS,
    GRID_TABLE_FACTS,
    Adapter,
    Mastaba,
    register_adapter,
)

#: A table on a slide is a grid, and it is described in the same words a grid is
#: described in everywhere else — so the same recognition rules can read it. The
#: words themselves come from the contract, not from a copy kept here (K2.22).
CELL_FACTS = GRID_CELL_FACTS
TABLE_FACTS = GRID_TABLE_FACTS

#: The vocabularies of what only a deck has (K2.19).
SLIDE_FACTS = ("shape_count",)
SHAPE_FACTS = ("shape_type", "is_title", "on_slide")
NOTES_FACTS = ("on_slide",)


def _described(image, name: str):
    """One property of a picture, or None where the header will not parse.

    Every descriptive property of this library's image object parses the
    header on access -- `content_type`, `size` and `ext` alike -- so each is
    a place an unparseable picture can raise. The bytes are reachable
    regardless, and a property that cannot be read is emitted as unknown.
    """
    try:
        return getattr(image, name)
    except Exception:
        return None



class PptxAdapter(Adapter):
    """Read a deck into slide, shape and notes observations."""

    format = "pptx"
    version = "0.5.0"          # 0.4.0: a vocabulary declared per record kind
    skip_reasons = PART_SKIPS
    extensions = (".pptx",)
    fact_sets = {
        "figure": FIGURE_FACTS,
        "slide": SLIDE_FACTS,
        "shape": SHAPE_FACTS,
        "notes": NOTES_FACTS,
        "table": TABLE_FACTS,
        "cell": CELL_FACTS,
    }

    def _figure(self, shape, number: int):
        """A picture on a slide, with the box the format states for it."""
        if self.store is None:
            return None
        # Reaching the bytes and describing them are two different operations
        # here, and only the second parses. A picture whose header will not parse
        # keeps its bytes and its placement, and loses only what the header would
        # have said; a picture whose part is gone has nothing to keep (K6.9).
        try:
            image = shape.image
            payload = image.blob
        except Exception:
            self._skip("image-part-unreadable")
            return None
        if not payload:
            self._skip("image-part-empty")
            return None
        media = _described(image, "content_type") or "application/octet-stream"
        pixels = _described(image, "size") or (None, None)
        digest = self.store.put(payload, media,
                                reference={"slide": number,
                                           "shape_id": int(shape.shape_id)})
        return Mastaba(
            kind="figure",
            locator=PptxLocator(slide=number, shape_id=int(shape.shape_id)),
            facts={
                "asset": digest,
                "source": "part",
                "media_type": media,
                "width": pixels[0],
                "height": pixels[1],
                "x": _points(shape.left),
                "y": _points(shape.top),
                "w": _points(shape.width),
                "h": _points(shape.height),
                "colorspace": None,
                "bits": None,
                "smask": None,
            },
        )

    def read(self, path: Path) -> Iterator[Mastaba]:
        from pptx import Presentation

        deck = Presentation(path)
        for number, slide in enumerate(deck.slides, start=1):
            title = slide.shapes.title
            # Identity is not reliable here: the library hands back a fresh
            # proxy object on each access, so `shape is title` is False even
            # for the title itself. The shape id is stable; compare that.
            title_id = None if title is None else title.shape_id
            yield Mastaba(
                kind="slide",
                locator=PptxLocator(slide=number),
                text=title.text if title is not None else None,
                facts={"shape_count": len(slide.shapes)},
            )

            for index, shape in enumerate(slide.shapes):
                if getattr(shape, "has_table", False):
                    yield from self._table(shape, number, index)
                    continue
                figure = self._figure(shape, number)
                if figure is not None:
                    yield figure
                    continue
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text
                if not text.strip():
                    continue
                yield Mastaba(
                    kind="shape",
                    locator=PptxLocator(slide=number, shape=index),
                    text=text,
                    facts={
                        "is_title": shape.shape_id == title_id,
                        "shape_type": str(shape.shape_type),
                        "on_slide": True,
                    },
                )

            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
                if notes.strip():
                    yield Mastaba(
                        kind="notes",
                        locator=PptxLocator(slide=number, notes=True),
                        text=notes,
                        facts={"on_slide": False},
                    )


    # ------------------------------------------------------------------ table

    def _table(self, shape, slide: int, index: int):
        """A slide table, as a resolved grid in the shared vocabulary."""
        table = shape.table
        rows = len(table.rows)
        cols = len(table.columns)

        yield Mastaba(
            kind="table",
            locator=PptxLocator(slide=slide, shape=index),
            facts={"n_rows": rows, "n_grid_cols": cols, "style": None, "ragged": False},
        )

        # Where each covered position's merge began, so a position covered from
        # ABOVE can be told from one covered from the LEFT.
        origin_of = {}
        for row in range(rows):
            for col in range(cols):
                cell = table.cell(row, col)
                if not cell.is_merge_origin:
                    continue
                for r in range(row, row + cell.span_height):
                    for c in range(col, col + cell.span_width):
                        origin_of[(r, c)] = (row, col)

        for row in range(rows):
            for col in range(cols):
                cell = table.cell(row, col)
                if cell.is_spanned:
                    origin = origin_of.get((row, col))
                    if origin is None or origin[1] != col:
                        # Covered from the left: the span already says so, and a
                        # word-processing file does not store this position either.
                        continue
                    # Covered from above. A word-processing reader emits this as a
                    # continuation, so this one does too — the shared mapping walks
                    # continuations to find how far a merge reaches, and a format
                    # that stayed silent here would lose the extent.
                    yield Mastaba(
                        kind="cell",
                        locator=PptxLocator(slide=slide, shape=index, row=row, col=col),
                        facts={"span": 1, "vmerge": "continue", "empty": True,
                               "fillable": False, "marker": None, "bold": False,
                               "shaded": False},
                    )
                    continue
                text = cell.text.strip()
                yield Mastaba(
                    kind="cell",
                    locator=PptxLocator(slide=slide, shape=index, row=row, col=col),
                    text=text or None,
                    facts={
                        "span": cell.span_width,
                        "vmerge": "restart" if cell.span_height > 1 else None,
                        "empty": not text,
                        "fillable": False,
                        "marker": None,
                        "bold": self._bold(cell),
                        "shaded": self._shaded(cell),
                    },
                )

    @staticmethod
    def _bold(cell) -> bool:
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.bold:
                    return True
        return False

    @staticmethod
    def _shaded(cell) -> bool:
        try:
            return str(cell.fill.type) not in ("None", "MSO_FILL_TYPE.BACKGROUND", "-1")
        except Exception:
            return False


register_adapter(PptxAdapter())


#: A point is 12700 English Metric Units. The presentation format states where a
#: picture landed, so this reader fills the geometry the word-processing one
#: cannot: what a format says, its reader reports.
_EMU_PER_POINT = 12700


def _points(value):
    return round(value / _EMU_PER_POINT, 2) if value is not None else None
